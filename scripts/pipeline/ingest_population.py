"""Ingesta de población distrital proyectada del INEI (base para tasas por 100 000).

Fuente: "Perú: Población Total Proyectada al 30 de Junio de cada año, según
Departamento, Provincia y Distrito, 2018-2026" (INEI, XLSX). Es la vía que el
proyecto dejó pendiente en `data/README.md` para normalizar los incidentes por
población.

Se extraen los distritos de Lima Metropolitana y Callao (ubigeo 1501/0701) y se
emite una serie larga (ubigeo, distrito, año, población) que se une por `ubigeo`.

Uso (desde la raíz del repositorio):

    python scripts/pipeline/ingest_population.py --refresh

Salidas:
    data/raw/population/poblacion_distrital_2018_2026.xlsx   (crudo, no versionado)
    data/processed/population_lima.csv                       (ubigeo, año, población)
    data/processed/population_quality_report.json
"""

from __future__ import annotations

import argparse
import json
import sys
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PIPELINE_DIR = Path(__file__).resolve().parent
REPO_ROOT = PIPELINE_DIR.parents[1]

SOURCE_URL = (
    "https://cdn.www.gob.pe/uploads/document/file/8261096/"
    "6894980-peru-poblacion-total-proyectada-al-30-de-junio-de-cada-ano-segun-"
    "departamento-provincia-y-distrito-2018-2026.xlsx?v=1768402069"
)
USER_AGENT = "RiskmapPipeline/1.0 (ingesta población INEI; proyecto academico)"
DEFAULT_RAW = REPO_ROOT / "data" / "raw" / "population" / "poblacion_distrital_2018_2026.xlsx"
SCOPE_PREFIXES = ("1501", "0701")  # Lima Metropolitana + Callao

COLUMNS = ["ubigeo", "name", "district", "year", "population", "source"]


def download(dest: Path) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(SOURCE_URL, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response:
        dest.write_bytes(response.read())


def _is_district(ubigeo: str) -> bool:
    """Distrito = ubigeo de 6 dígitos cuyo sufijo no es '00' (provincia/departamento)."""
    return len(ubigeo) == 6 and ubigeo.isdigit() and ubigeo[4:] != "00"


def parse_population(path: Path) -> tuple[list[dict], dict]:
    """Devuelve (registros en ámbito, stats) desde el XLSX del INEI."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    header_index = next(
        (i for i, row in enumerate(rows) if row and str(row[0] or "").strip().upper() == "UBIGEO"),
        None,
    )
    if header_index is None:
        raise SystemExit("No se encontró la fila de encabezado (UBIGEO) en el XLSX de población.")

    # La fila de años es la siguiente que tenga enteros en las columnas >= 2.
    year_row = None
    for row in rows[header_index + 1:]:
        years = [
            (column, int(cell))
            for column, cell in enumerate(row)
            if column >= 2 and isinstance(cell, (int, float)) and 2000 <= cell <= 2100
        ]
        if years:
            year_row = years
            break
    if not year_row:
        raise SystemExit("No se encontró la fila de años en el XLSX de población.")

    data_start = header_index + 2
    records: list[dict] = []
    stats = {"rows_raw": 0, "fuera_de_ambito": 0, "no_distrito": 0}
    for row in rows[data_start:]:
        if not row or row[0] in (None, ""):
            continue
        ubigeo = str(row[0]).strip().split(".")[0].zfill(6)
        stats["rows_raw"] += 1
        if not ubigeo.startswith(SCOPE_PREFIXES):
            stats["fuera_de_ambito"] += 1
            continue
        if not _is_district(ubigeo):
            stats["no_distrito"] += 1
            continue
        name = str(row[1] or "").strip()
        district = name.split(",")[-1].strip() if name else ubigeo
        for column, year in year_row:
            value = row[column] if column < len(row) else None
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                continue
            records.append(
                {
                    "ubigeo": ubigeo,
                    "name": name,
                    "district": district.upper(),
                    "year": year,
                    "population": int(value),
                    "source": "INEI",
                }
            )
    return records, stats


def main() -> int:
    parser = argparse.ArgumentParser(description="Población distrital proyectada (INEI) para Lima/Callao.")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_RAW)
    parser.add_argument("--refresh", action="store_true", help="Descargar el XLSX de nuevo.")
    args = parser.parse_args()

    if args.refresh or not args.xlsx.exists():
        print(f"Descargando población distrital a {args.xlsx} ...")
        download(args.xlsx)

    records, stats = parse_population(args.xlsx)
    frame = pd.DataFrame(records, columns=COLUMNS)
    if not frame.empty:
        frame = frame.sort_values(["year", "ubigeo"]).reset_index(drop=True)

    out_dir = REPO_ROOT / "data" / "processed"
    out_dir.mkdir(parents=True, exist_ok=True)
    dataset_path = out_dir / "population_lima.csv"
    report_path = out_dir / "population_quality_report.json"

    frame.to_csv(dataset_path, index=False, encoding="utf-8")
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": SOURCE_URL,
        "scope_prefixes": list(SCOPE_PREFIXES),
        "rows": int(len(frame)),
        "districts": int(frame["ubigeo"].nunique()) if not frame.empty else 0,
        "years": sorted(int(y) for y in frame["year"].unique()) if not frame.empty else [],
        "latest_year_population": int(
            frame[frame["year"] == frame["year"].max()]["population"].sum()
        )
        if not frame.empty
        else 0,
        **stats,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"Población: {dataset_path} ({len(frame)} filas)")
    print(f"Reporte:   {report_path}")
    print(
        f"Distritos: {report['districts']} | años: {report['years'][:1]}–{report['years'][-1:]} "
        f"| población {report['years'][-1]}: {report['latest_year_population']:,}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
